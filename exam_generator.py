import os
import tempfile
from tkinter import *
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

# --------------------- إعدادات الخط العربي ---------------------
FONT_PATH = "Amiri-Regular.ttf"
try:
    pdfmetrics.registerFont(TTFont('ArabicFont', FONT_PATH))
    ARABIC_FONT = 'ArabicFont'
    print("✅ تم تحميل خط أميري")
except:
    ARABIC_FONT = 'Helvetica'
    print("⚠️ خط عربي غير موجود، سيستخدم الافتراضي")

# --------------------- الوظيفة الرئيسية ---------------------
def generate_papers(excel_path, output_dir):
    """
    توليد أوراق الأسئلة لكل طالب ولكل مادة مسجل بها.
    - كود المادة = رقم العمود (1 للعمود E، 2 للعمود F، ...)
    - المربعات الثلاثة في الأسفل: [مربع الدرجة (أزرق)] [QR] [كود المادة]
    - أبعاد المربع: 40×40 ملم
    """
    # قراءة ملف الإكسيل
    wb = load_workbook(excel_path, data_only=True)
    sheet = wb.active

    # استخراج رؤوس المواد من الأعمدة E إلى S مع أرقامها
    material_info = {}  # {column_index: {'code': number, 'name': header}}
    for col in range(5, 20):  # E=5, S=19
        header = sheet.cell(1, col).value
        if header and str(header).strip():
            code = col - 4  # E=1, F=2, G=3, ...
            material_info[col] = {'code': code, 'name': str(header).strip()}

    if not material_info:
        messagebox.showerror("خطأ", "لم يتم العثور على أعمدة مواد (E إلى S)")
        return

    # إنشاء مجلد الإخراج
    os.makedirs(output_dir, exist_ok=True)

    # مسار مجلد QR (بجانب ملف الإكسيل)
    base_dir = os.path.dirname(excel_path)
    qr_folder = os.path.join(base_dir, "qr_pict")

    total_pages = 0

    # التكرار على صفوف الطلاب (من الصف الثاني)
    for row in range(2, sheet.max_row + 1):
        student_id = str(sheet.cell(row, 1).value or '').strip()
        student_name = str(sheet.cell(row, 2).value or '').strip()
        student_class = str(sheet.cell(row, 3).value or '').strip()

        if not student_id or not student_name:
            continue

        # التكرار على أعمدة المواد
        for col, info in material_info.items():
            material_code = info['code']      # رقم المادة (1, 2, 3, ...)
            material_name = info['name']      # اسم المادة (نص)

            # التحقق مما إذا كان الطالب مسجلاً في هذه المادة
            grade_cell = sheet.cell(row, col).value
            if grade_cell is None:
                continue

            # إنشاء ملف PDF
            pdf_path = os.path.join(output_dir, f"{student_id}_M{material_code}.pdf")
            c = canvas.Canvas(pdf_path, pagesize=A4)
            width, height = A4

            # ---- أعلى الصفحة ----
            # جهة اليمين: اسم الطالب ورقم الجلوس
            c.setFont(ARABIC_FONT, 18)
            c.drawRightString(width - 50, height - 50, f"{student_name}  |  {student_id}")

            # جهة اليسار: اسم المادة (للتذكير)
            c.setFont(ARABIC_FONT, 14)
            c.drawString(50, height - 50, f"المادة: {material_name}")

            # ---- أسفل الصفحة: ثلاثة مربعات متجاورة (40×40 ملم) ----
            box_size = 40 * mm       # 40 ملم
            spacing = 10 * mm        # المسافة بين المربعات
            y_start = 20 * mm        # المسافة من أسفل الصفحة

            # المربع الأول (الأيسر): مربع أزرق فاتح فارغ للدرجة
            x1 = (width - (box_size * 3 + spacing * 2)) / 2  # توسيط المجموعة
            c.setFillColorRGB(0.9, 0.95, 1.0)  # أزرق فاتح جداً
            c.rect(x1, y_start, box_size, box_size, stroke=1, fill=1)
            c.setFillColorRGB(0, 0, 0)  # إعادة اللون للأسود

            # المربع الثاني (الوسط): صورة QR
            x2 = x1 + box_size + spacing
            c.rect(x2, y_start, box_size, box_size, stroke=1, fill=0)
            qr_path = os.path.join(qr_folder, f"{student_id}.png")
            if os.path.exists(qr_path):
                try:
                    img = Image.open(qr_path)
                    with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as tmp:
                        img.save(tmp.name)
                        # ترك هامش 5 ملم داخل المربع
                        margin = 5 * mm
                        c.drawImage(tmp.name,
                                    x2 + margin, y_start + margin,
                                    width=box_size - 2*margin,
                                    height=box_size - 2*margin,
                                    preserveAspectRatio=True)
                        os.unlink(tmp.name)
                except Exception as e:
                    print(f"⚠️ فشل تحميل QR للطالب {student_id}: {e}")
            else:
                c.setFont('Helvetica', 10)
                c.drawCentredString(x2 + box_size/2, y_start + box_size/2 - 5, "QR غير موجود")

            # المربع الثالث (الأيمن): كود المادة (الرقم فقط)
            x3 = x2 + box_size + spacing
            c.rect(x3, y_start, box_size, box_size, stroke=1, fill=0)
            c.setFont(ARABIC_FONT, 20)
            c.drawCentredString(x3 + box_size/2, y_start + box_size/2 - 5, str(material_code))

            # حفظ الصفحة
            c.save()
            total_pages += 1
            print(f"✅ تم توليد: {pdf_path}")

    messagebox.showinfo("نجاح", f"تم توليد {total_pages} ورقة بنجاح في:\n{output_dir}")

# --------------------- الواجهة الرسومية ---------------------
class App:
    def __init__(self, root):
        self.root = root
        root.title("مولد أوراق الأسئلة - مربعات 40×40 ملم")
        root.geometry("500x250")
        root.configure(bg='#f0f0f0')

        self.excel_path = StringVar()

        main_frame = Frame(root, bg='#f0f0f0', padx=20, pady=20)
        main_frame.pack(fill=BOTH, expand=True)

        Label(main_frame, text="ملف الإكسيل:", bg='#f0f0f0', font=('Arial', 12)).grid(row=0, column=0, sticky='w')
        Entry(main_frame, textvariable=self.excel_path, width=35).grid(row=0, column=1, padx=5)
        Button(main_frame, text="استعراض", command=self.browse_excel).grid(row=0, column=2)

        Button(main_frame, text="توليد الأوراق لجميع الطلاب والمواد",
               command=self.generate_all, bg='#4CAF50', fg='white',
               font=('Arial', 14), padx=20, pady=10).grid(row=2, column=1, pady=30)

        self.log_text = Text(main_frame, height=4, state='disabled', bg='white')
        self.log_text.grid(row=3, column=0, columnspan=3, sticky='we', pady=10)

    def browse_excel(self):
        filename = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if filename:
            self.excel_path.set(filename)

    def generate_all(self):
        path = self.excel_path.get()
        if not path or not os.path.exists(path):
            messagebox.showwarning("تنبيه", "يرجى اختيار ملف الإكسيل أولاً")
            return

        output_dir = filedialog.askdirectory(title="اختر مجلد حفظ الأوراق")
        if not output_dir:
            return

        self.log("⏳ جاري التوليد...")
        self.root.update()

        try:
            generate_papers(path, output_dir)
            self.log("✅ تم الانتهاء بنجاح")
        except Exception as e:
            self.log(f"❌ خطأ: {e}")
            messagebox.showerror("خطأ", str(e))

    def log(self, msg):
        self.log_text.config(state='normal')
        self.log_text.insert(END, msg + "\n")
        self.log_text.see(END)
        self.log_text.config(state='disabled')

# --------------------- تشغيل التطبيق ---------------------
if __name__ == "__main__":
    root = Tk()
    app = App(root)
    root.mainloop()
